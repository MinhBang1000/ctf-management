import io

from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.models.report import Report


class ExportNotAvailableError(Exception):
    """Report has no structured data to export (e.g. a weekly report, or
    a semester report generated before Phase 5's `data` column existed)."""


TABLE_STYLE = TableStyle(
    [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0d7d8f")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cccccc")),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f7f9f9")]),
    ]
)


def _require_data(report: Report) -> dict:
    if report.type != "semester" or not report.data:
        raise ExportNotAvailableError("Only semester reports have exportable structured data")
    return report.data


def render_semester_report_pdf(report: Report) -> bytes:
    data = _require_data(report)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter)
    styles = getSampleStyleSheet()
    elements = [
        Paragraph(f"Semester Report — {data['semester_name']}", styles["Title"]),
        Paragraph(f"{data['period_start']} to {data['period_end']}", styles["Normal"]),
        Spacer(1, 12),
    ]

    if data["spans_platform_switch"]:
        names = ", ".join(p["name"] for p in data["platforms_used"])
        elements.append(
            Paragraph(
                f"<i>⚠ This semester spans a Platform switch: {names}. "
                "Completion and points are labeled per-platform below.</i>",
                styles["Normal"],
            )
        )
        elements.append(Spacer(1, 12))

    elements.append(Paragraph("Weekly Trend", styles["Heading2"]))
    trend_rows = [["Week", "Platform(s)", "Challenges", "Completion %"]]
    for w in data["weekly_trend"]:
        trend_rows.append(
            [str(w["week_number"]), ", ".join(w["platform_names"]), str(w["challenge_count"]), f"{w['completion_pct']}%"]
        )
    trend_table = Table(trend_rows, hAlign="LEFT")
    trend_table.setStyle(TABLE_STYLE)
    elements.append(trend_table)
    elements.append(Spacer(1, 20))

    elements.append(Paragraph("Per-Member Ranking", styles["Heading2"]))
    member_rows = [["Rank", "Member", "Points", "Done", "Late", "Missing"]]
    for m in data["members"]:
        member_rows.append(
            [str(m["rank"]), m["member_name"], str(m["cumulative_points"]), str(m["done_count"]), str(m["late_count"]), str(m["missing_count"])]
        )
    member_table = Table(member_rows, hAlign="LEFT")
    member_table.setStyle(TABLE_STYLE)
    elements.append(member_table)

    if data["spans_platform_switch"]:
        elements.append(Spacer(1, 20))
        elements.append(Paragraph("Per-Member, Per-Platform Breakdown", styles["Heading2"]))
        bp_rows = [["Member", "Platform", "Points", "Done", "Late", "Missing"]]
        for m in data["members"]:
            for bp in m["by_platform"]:
                bp_rows.append(
                    [m["member_name"], bp["platform_name"], str(bp["points"]), str(bp["done"]), str(bp["late"]), str(bp["missing"])]
                )
        bp_table = Table(bp_rows, hAlign="LEFT")
        bp_table.setStyle(TABLE_STYLE)
        elements.append(bp_table)

    doc.build(elements)
    return buffer.getvalue()


def render_semester_report_excel(report: Report) -> bytes:
    data = _require_data(report)
    wb = Workbook()

    ws = wb.active
    ws.title = "Member Ranking"
    ws.append(["Rank", "Member", "Points", "Done", "Late", "Missing"])
    for m in data["members"]:
        ws.append([m["rank"], m["member_name"], m["cumulative_points"], m["done_count"], m["late_count"], m["missing_count"]])

    ws2 = wb.create_sheet("Weekly Trend")
    ws2.append(["Week", "Platform(s)", "Challenges", "Completion %"])
    for w in data["weekly_trend"]:
        ws2.append([w["week_number"], ", ".join(w["platform_names"]), w["challenge_count"], w["completion_pct"]])

    if data["spans_platform_switch"]:
        ws3 = wb.create_sheet("Per-Platform Breakdown")
        ws3.append(["Member", "Platform", "Points", "Done", "Late", "Missing"])
        for m in data["members"]:
            for bp in m["by_platform"]:
                ws3.append([m["member_name"], bp["platform_name"], bp["points"], bp["done"], bp["late"], bp["missing"]])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
